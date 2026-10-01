#!/usr/bin/env python3
"""2026년 주차별 일별+주간합계 리포트 엑셀 추출 (샘플수/청첩장주문수/청첩장매출액, 브랜드별, 전년 동요일 비교 포함)"""
import sys
import importlib.util
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

_spec = importlib.util.spec_from_file_location(
    "daily_report_dashboard", r"C:\src\barunson-database-reference\user\daily-report-dashboard.py")
_mod = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_mod)
get_conn, BRANDS, _brand_key, ALL_GUBUN_SQL = _mod.get_conn, _mod.BRANDS, _mod._brand_key, _mod.ALL_GUBUN_SQL

START = date(2026, 1, 4)          # 2026년 2주차 시작
END = date.today()                 # 오늘까지 (이번 주는 미완결)
WEEK1_START = date(2025, 12, 28)   # 2026년 1주차 시작(일) 기준점
YOY_OFFSET = timedelta(weeks=52)   # 항상 동일 요일

WEEKDAY_KO = ['월', '화', '수', '목', '금', '토', '일']


def week_number(d):
    wstart = d - timedelta(days=(d.weekday() + 1) % 7)  # 그 주의 일요일
    return ((wstart - WEEK1_START).days // 7) + 1


def _bucket(out, d, bk):
    return out.setdefault(d, {}).setdefault(bk, {'total': 0, 'member': 0, 'guest': 0})


def fetch_sample_range(conn, s, e):
    cur = conn.cursor(as_dict=True)
    e1 = (e + timedelta(days=1)).strftime('%Y-%m-%d')
    cur.execute(f"""
        SELECT CAST(REQUEST_DATE AS DATE) as d, SALES_GUBUN as sg,
            CASE WHEN MEMBER_ID IS NOT NULL AND MEMBER_ID <> '' THEN 'member' ELSE 'guest' END as ut,
            COUNT(*) as cnt
        FROM CUSTOM_SAMPLE_ORDER
        WHERE REQUEST_DATE >= %s AND REQUEST_DATE < %s AND STATUS_SEQ >= 1
          AND SALES_GUBUN IN {ALL_GUBUN_SQL}
        GROUP BY CAST(REQUEST_DATE AS DATE), SALES_GUBUN,
            CASE WHEN MEMBER_ID IS NOT NULL AND MEMBER_ID <> '' THEN 'member' ELSE 'guest' END
    """, (s.strftime('%Y-%m-%d'), e1))
    out = {}
    for row in cur.fetchall():
        bk = _brand_key(row['sg'])
        if not bk:
            continue
        b = _bucket(out, row['d'], bk)
        b[row['ut']] += row['cnt']
        b['total'] += row['cnt']
    return out


def fetch_settle_count_range(conn, s, e):
    cur = conn.cursor(as_dict=True)
    e1 = (e + timedelta(days=1)).strftime('%Y-%m-%d')
    cur.execute(f"""
        SELECT CAST(settle_date AS DATE) as d, sales_Gubun as sg,
            CASE WHEN member_id IS NOT NULL AND member_id <> '' THEN 'member' ELSE 'guest' END as ut,
            COUNT(DISTINCT order_seq) as cnt
        FROM custom_order
        WHERE settle_date >= %s AND settle_date < %s AND sales_Gubun IN {ALL_GUBUN_SQL}
        GROUP BY CAST(settle_date AS DATE), sales_Gubun,
            CASE WHEN member_id IS NOT NULL AND member_id <> '' THEN 'member' ELSE 'guest' END
    """, (s.strftime('%Y-%m-%d'), e1))
    out = {}
    for row in cur.fetchall():
        bk = _brand_key(row['sg'])
        if not bk:
            continue
        b = _bucket(out, row['d'], bk)
        b[row['ut']] += row['cnt']
        b['total'] += row['cnt']
    return out


def fetch_settle_revenue_range(conn, s, e):
    cur = conn.cursor(as_dict=True)
    e1 = (e + timedelta(days=1)).strftime('%Y-%m-%d')
    cur.execute(f"""
        SELECT CAST(settle_date AS DATE) as d, sales_Gubun as sg,
            CASE WHEN member_id IS NOT NULL AND member_id <> '' THEN 'member' ELSE 'guest' END as ut,
            SUM(settle_price) as rev
        FROM custom_order
        WHERE settle_date >= %s AND settle_date < %s AND sales_Gubun IN {ALL_GUBUN_SQL}
        GROUP BY CAST(settle_date AS DATE), sales_Gubun,
            CASE WHEN member_id IS NOT NULL AND member_id <> '' THEN 'member' ELSE 'guest' END
    """, (s.strftime('%Y-%m-%d'), e1))
    out = {}
    for row in cur.fetchall():
        bk = _brand_key(row['sg'])
        if not bk:
            continue
        b = _bucket(out, row['d'], bk)
        v = row['rev'] or 0
        b[row['ut']] += v
        b['total'] += v
    return out


def get(d_map, d, bk, skey='total'):
    return d_map.get(d, {}).get(bk, {}).get(skey, 0)


def main():
    conn = get_conn()
    yoy_start, yoy_end = START - YOY_OFFSET, END - YOY_OFFSET

    print(f"기간: {START} ~ {END} / 전년: {yoy_start} ~ {yoy_end}", file=sys.stderr)

    sample_cur = fetch_sample_range(conn, START, END)
    sample_yoy = fetch_sample_range(conn, yoy_start, yoy_end)
    cnt_cur = fetch_settle_count_range(conn, START, END)
    cnt_yoy = fetch_settle_count_range(conn, yoy_start, yoy_end)
    rev_cur = fetch_settle_revenue_range(conn, START, END)
    rev_yoy = fetch_settle_revenue_range(conn, yoy_start, yoy_end)
    conn.close()

    brand_keys = [k for k, _, _ in BRANDS]
    brand_names = {k: n for k, n, _ in BRANDS}
    metrics = [
        ('샘플수', sample_cur, sample_yoy),
        ('청첩장주문수', cnt_cur, cnt_yoy),
        ('청첩장매출액', rev_cur, rev_yoy),
    ]

    # 행 정의: (라벨, 조회맵, 브랜드키, 구분키(total/member/guest), 전년여부)
    SUBTYPES = [('전체', 'total'), ('회원', 'member'), ('비회원', 'guest')]
    row_defs = []
    for mname, cur_map, yoy_map in metrics:
        for bk in brand_keys:
            for slabel, skey in SUBTYPES:
                row_defs.append((f"{mname}_{brand_names[bk]}_{slabel}", cur_map, bk, skey, False))
                row_defs.append((f"{mname}_{brand_names[bk]}_{slabel}_전년", yoy_map, bk, skey, True))

    # 날짜 목록 + 주차별 그룹
    dates = []
    d = START
    while d <= END:
        dates.append(d)
        d += timedelta(days=1)

    weeks = []  # [(주차번호, [날짜,...]), ...]
    for d in dates:
        wk = week_number(d)
        if not weeks or weeks[-1][0] != wk:
            weeks.append((wk, []))
        weeks[-1][1].append(d)

    wb = Workbook()
    ws = wb.active
    ws.title = "주차별리포트"

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    sum_fill = PatternFill("solid", fgColor="FFF2CC")

    lbl = ws.cell(row=1, column=1, value="지표")
    lbl.font = Font(bold=True)
    lbl.fill = header_fill

    date_col = {}
    week_sum_col = {}
    col = 2
    for wk, wdates in weeks:
        for d in wdates:
            c = ws.cell(row=1, column=col, value=f"{d.strftime('%Y-%m-%d')}({WEEKDAY_KO[d.weekday()]})")
            c.font = Font(bold=True)
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center")
            date_col[d] = col
            col += 1
        c = ws.cell(row=1, column=col, value=f"{wk}주차 합계")
        c.font = Font(bold=True)
        c.fill = sum_fill
        c.alignment = Alignment(horizontal="center")
        week_sum_col[wk] = col
        col += 1
    last_col = col - 1

    for r_i, (label, val_map, bk, skey, is_yoy) in enumerate(row_defs, start=2):
        lc = ws.cell(row=r_i, column=1, value=label)
        lc.font = Font(bold=True)
        for wk, wdates in weeks:
            wsum = 0
            for d in wdates:
                src_d = (d - YOY_OFFSET) if is_yoy else d
                v = get(val_map, src_d, bk, skey)
                ws.cell(row=r_i, column=date_col[d], value=v)
                wsum += v
            sc = ws.cell(row=r_i, column=week_sum_col[wk], value=wsum)
            sc.font = Font(bold=True)
            sc.fill = sum_fill

    ws.freeze_panes = "B2"
    ws.column_dimensions['A'].width = 24
    for c in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 12

    out_path = r"C:\Users\LG\Desktop\2026_주차별_리포트_가로형_회원구분.xlsx"
    wb.save(out_path)
    print(f"저장 완료: {out_path}", file=sys.stderr)


if __name__ == '__main__':
    main()
